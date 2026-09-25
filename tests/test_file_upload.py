import io
import json
from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import pytest
from flask.testing import FlaskClient

from app import create_app
from db import create_user, get_session, init_db
from services.builder import advance_session, build_system_prompt, create_session_with_idea
from services.extractor import (
    ExtractionError,
    cap_words,
    count_words,
    extract_text_from_file,
    process_attachment,
    summarize_document,
)


@pytest.fixture
def test_db(tmp_path: Path) -> Path:
    """Fixture to create a fresh test database."""
    db_file = tmp_path / "test_upload.db"
    init_db(db_file)
    return db_file


@pytest.fixture
def client(tmp_path: Path) -> FlaskClient:
    """Fixture providing a test client configured with an isolated test database."""
    test_db = tmp_path / "test_api_upload.db"
    test_app = create_app(db_path=test_db)
    test_app.config["TESTING"] = True
    test_app.config["RATELIMIT_ENABLED"] = False
    test_app.config["CSRF_ENABLED"] = False
    user_id = create_user("testuser", "mock_hash", db_path=test_db)
    test_client = test_app.test_client()
    with test_client.session_transaction() as sess:
        sess["authenticated"] = True
        sess["user_id"] = user_id
        sess["username"] = "testuser"
    return test_client


def make_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    """Helper to generate in-memory docx bytes."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r_idx, row in enumerate(table_rows):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- Extraction Tests ---

def test_extract_text_txt() -> None:
    """Verify extracting raw text from a .txt file."""
    content = "This is a plain text file for prompt context."
    extracted = extract_text_from_file(content.encode("utf-8"), "notes.txt")
    assert extracted == content


def test_extract_text_pdf() -> None:
    """Verify extracting raw text from a valid .pdf file."""
    sample_pdf_path = Path(__file__).resolve().parent / "sample_math_syllabus.pdf"
    assert sample_pdf_path.exists(), "Sample PDF file should exist"
    with open(sample_pdf_path, "rb") as f:
        pdf_bytes = f.read()

    extracted = extract_text_from_file(pdf_bytes, "sample_math_syllabus.pdf")
    assert "Calculus" in extracted
    assert "Integrals" in extracted


def test_extract_text_docx() -> None:
    """Verify extracting raw text from a .docx file including paragraphs and tables."""
    docx_bytes = make_docx_bytes(
        paragraphs=["Module 1: Introduction to AI.", "Module 2: Neural Networks."],
        table_rows=[["Topic", "Weeks"], ["Backpropagation", "2 weeks"]],
    )
    extracted = extract_text_from_file(docx_bytes, "course.docx")
    assert "Module 1: Introduction to AI." in extracted
    assert "Module 2: Neural Networks." in extracted
    assert "Backpropagation | 2 weeks" in extracted


def test_extract_text_word_cap() -> None:
    """Verify that extraction caps content at roughly 15,000 words."""
    words = [f"word{i}" for i in range(16000)]
    content = " ".join(words)
    extracted = extract_text_from_file(content.encode("utf-8"), "long.txt", max_words=15000)
    assert count_words(extracted) == 15000


def test_extract_text_corrupt_pdf() -> None:
    """Verify that a corrupted PDF file raises ExtractionError."""
    corrupt_bytes = b"This is definitely not a valid PDF header or content."
    with pytest.raises(ExtractionError, match="Failed to extract text from PDF"):
        extract_text_from_file(corrupt_bytes, "corrupt.pdf")


def test_extract_text_corrupt_docx() -> None:
    """Verify that a corrupted DOCX file raises ExtractionError."""
    corrupt_bytes = b"Random binary gibberish that is not a zip archive"
    with pytest.raises(ExtractionError, match="Failed to extract text from DOCX"):
        extract_text_from_file(corrupt_bytes, "corrupt.docx")


def test_extract_text_unsupported_format() -> None:
    """Verify that unsupported file formats raise ExtractionError."""
    with pytest.raises(ExtractionError, match="Unsupported file format"):
        extract_text_from_file(b"print('hello')", "script.py")

    with pytest.raises(ExtractionError, match="Unsupported file format"):
        extract_text_from_file(b"\x89PNG\r\n\x1a\n", "image.png")


def test_extract_text_empty_file() -> None:
    """Verify that empty file bytes raise ExtractionError."""
    with pytest.raises(ExtractionError, match="Uploaded file is empty"):
        extract_text_from_file(b"", "empty.txt")


# --- Short vs Long Branching Tests ---

def test_short_attachment_no_summarization() -> None:
    """Verify that short attachments under 3,000 words are returned directly without calling summarization."""
    short_content = "Short document context with only a few words."
    with patch("services.extractor.summarize_document") as mock_summarize:
        result_text, was_summarized = process_attachment(short_content.encode("utf-8"), "short.txt")
        assert was_summarized is False
        assert result_text == short_content
        mock_summarize.assert_not_called()


def test_long_attachment_invokes_summarization() -> None:
    """Verify that long attachments over 3,000 words trigger plain summarization via one LLM call."""
    long_words = [f"word{i}" for i in range(3500)]
    long_content = " ".join(long_words)
    expected_summary = "This is a concise summary of the 3500-word document."

    with patch("services.extractor.call_llm", return_value=expected_summary) as mock_llm:
        result_text, was_summarized = process_attachment(long_content.encode("utf-8"), "long.txt")
        assert was_summarized is True
        assert result_text == expected_summary
        mock_llm.assert_called_once()
        # Verify call_llm was invoked with plain json_mode=False and prompt from prompts/summarize.txt
        call_args, call_kwargs = mock_llm.call_args
        assert call_kwargs.get("json_mode") is False
        messages = call_args[0]
        assert "Summarize the following document" in messages[0]["content"]


# --- Graceful Degradation Tests ---

def test_flow_degrades_gracefully_on_corrupt_file(test_db: Path) -> None:
    """Verify session flow degrades gracefully and continues with user's idea if extraction fails."""
    mock_reply = json.dumps({
        "status": "ask",
        "type": "study",
        "question": "What level of calculus?",
    })
    corrupt_attachment = ("corrupted.pdf", b"Not a real PDF stream")

    with patch("services.builder.call_llm", return_value=mock_reply):
        result = create_session_with_idea(
            "Help me study calculus limits",
            attachment=corrupt_attachment,
            db_path=test_db,
        )
        assert result["status"] == "ask"
        assert result["has_attachment"] is False
        assert result["attachment_filename"] is None

        session = get_session(result["session_id"], db_path=test_db)
        assert session is not None
        assert session["has_attachment"] == 0
        assert session["attachment_filename"] is None
        assert session["attachment_text"] is None


def test_flow_degrades_gracefully_on_unsupported_file(test_db: Path) -> None:
    """Verify session flow degrades gracefully on unsupported file extensions."""
    mock_reply = json.dumps({
        "status": "ask",
        "type": "study",
        "question": "What level of calculus?",
    })
    unsupported_attachment = ("image.png", b"\x89PNG\r\n\x1a\n")

    with patch("services.builder.call_llm", return_value=mock_reply):
        result = create_session_with_idea(
            "Help me study calculus limits",
            attachment=unsupported_attachment,
            db_path=test_db,
        )
        assert result["status"] == "ask"
        assert result["has_attachment"] is False
        assert result["attachment_filename"] is None


# --- Context Injection & Multi-Turn Persistence Tests ---

def test_attachment_context_included_in_system_prompt() -> None:
    """Verify that build_system_prompt includes the required material prompt template."""
    prompt_with_material = build_system_prompt(
        "study",
        attachment_text="Vector calculus topics: Divergence and Curl.",
        attachment_filename="syllabus.txt",
    )
    assert "ATTACHED MATERIAL (syllabus.txt):" in prompt_with_material
    assert "Vector calculus topics: Divergence and Curl." in prompt_with_material
    assert "EMBED the condensed source material directly inline" in prompt_with_material


def test_attachment_context_short_vs_long_templates() -> None:
    """Verify short files use inline instructions and long files use paste placeholder instructions."""
    short_prompt = build_system_prompt(
        "study",
        attachment_text="Brief notes on calculus.",
        attachment_filename="notes.pdf",
        attachment_summarized=False,
    )
    assert "ATTACHED MATERIAL (notes.pdf):" in short_prompt
    assert "EMBED the condensed source material directly inline" in short_prompt
    assert "## Source Material" in short_prompt

    long_prompt = build_system_prompt(
        "study",
        attachment_text="Summary of comprehensive textbook.",
        attachment_filename="textbook.pdf",
        attachment_summarized=True,
    )
    assert "ATTACHED MATERIAL SUMMARY (textbook.pdf):" in long_prompt
    assert "[PASTE FULL textbook.pdf HERE]" in long_prompt


def test_attachment_context_persists_across_follow_ups(test_db: Path) -> None:
    """Verify that attachment text is persisted on the session and provided in system prompt on all turns."""
    mock_ask_turn1 = json.dumps({
        "status": "ask",
        "type": "study",
        "question": "Which specific theorem do you want to focus on?",
    })
    mock_ready_turn2 = json.dumps({
        "status": "ready",
        "type": "study",
        "final_prompt": "## Role\nMath Tutor\n## Task\nExplain Stokes Theorem based on provided syllabus.",
    })

    attachment = ("syllabus.txt", b"Syllabus: Stokes Theorem, Green Theorem, Gauss Divergence Theorem.")

    with patch("services.builder.call_llm", side_effect=[mock_ask_turn1, mock_ready_turn2]) as mock_llm:
        # Turn 1: Create session with attachment
        res1 = create_session_with_idea("Help me study calculus theorems", attachment=attachment, db_path=test_db)
        session_id = res1["session_id"]
        assert res1["has_attachment"] is True
        assert res1["attachment_filename"] == "syllabus.txt"

        # Check call_llm system prompt for Turn 1
        turn1_messages = mock_llm.call_args_list[0][0][0]
        turn1_system = turn1_messages[0]["content"]
        assert "Syllabus: Stokes Theorem" in turn1_system
        assert "ATTACHED MATERIAL (syllabus.txt):" in turn1_system

        # Turn 2: Follow-up answer
        res2 = advance_session(session_id, text="Stokes Theorem", db_path=test_db)
        assert res2["status"] == "ready"
        assert res2["has_attachment"] is True

        # Check call_llm system prompt for Turn 2: it MUST still contain the material
        turn2_messages = mock_llm.call_args_list[1][0][0]
        turn2_system = turn2_messages[0]["content"]
        assert "Syllabus: Stokes Theorem" in turn2_system
        assert "ATTACHED MATERIAL (syllabus.txt):" in turn2_system


def test_final_prompt_contains_identifiable_terms_from_attachment(test_db: Path) -> None:
    """Requirement 4: Verify generated final_prompt contains identifiable terms from the uploaded file."""
    sample_pdf_path = Path(__file__).resolve().parent / "sample_math_syllabus.pdf"
    assert sample_pdf_path.exists()
    pdf_bytes = sample_pdf_path.read_bytes()

    # The PDF contains: 'Math 201 Calculus III: Vector fields, Green Theorem, Stokes Theorem, and Surface Integrals.'
    mock_ready_reply = json.dumps({
        "status": "ready",
        "type": "study",
        "final_prompt": (
            "## Role\nAdvanced Calculus Tutor\n\n"
            "## Context\n"
            "The student is preparing for Math 201 Calculus III.\n\n"
            "## Source Material\n"
            "Math 201 Calculus III: Vector fields, Green Theorem, Stokes Theorem, and Surface Integrals.\n\n"
            "## Task\n"
            "Teach the relationship between Green's Theorem and Stokes' Theorem, and how they apply to Surface Integrals in Vector fields.\n"
        ),
    })

    with patch("services.builder.call_llm", return_value=mock_ready_reply) as mock_llm:
        res = create_session_with_idea(
            "Help me study for my upcoming exam",
            attachment=("sample_math_syllabus.pdf", pdf_bytes),
            db_path=test_db,
        )

        assert res["status"] == "ready"
        final_prompt = res["final_prompt"]

        # 1. Identifiable specific terms from the syllabus must be present
        assert "Stokes" in final_prompt or "Green" in final_prompt
        assert "Surface Integrals" in final_prompt
        assert "Vector fields" in final_prompt or "Calculus III" in final_prompt

        # 2. Short files must embed the source material inline under ## Source Material or ## Context
        assert "## Source Material" in final_prompt or "## Context" in final_prompt

        # 3. Must not contain generic placeholders like [TOPIC] or [PASTE YOUR NOTES / PDF TEXT HERE]
        assert "[TOPIC]" not in final_prompt
        assert "[PASTE YOUR NOTES / PDF TEXT HERE]" not in final_prompt

        # 4. Verify system prompt passed to call_llm contained the extracted PDF text
        call_system_prompt = mock_llm.call_args[0][0][0]["content"]
        assert "Green Theorem" in call_system_prompt
        assert "Stokes Theorem" in call_system_prompt
        assert "sample_math_syllabus.pdf" in call_system_prompt


def test_long_attachment_generates_paste_placeholder(test_db: Path) -> None:
    """Requirement 3(b): Verify long/summarized file uses [PASTE FULL <filename> HERE] placeholder."""
    mock_ready_reply = json.dumps({
        "status": "ready",
        "type": "study",
        "final_prompt": (
            "## Role\nBiochemistry Professor\n\n"
            "## Context\n"
            "Analyzing the attached full textbook chapter on cellular respiration.\n\n"
            "## Task\n"
            "Explain the electron transport chain and ATP synthesis based on the provided text.\n\n"
            "[PASTE FULL biochemistry_handbook.pdf HERE]\n"
        ),
    })

    with patch("services.extractor.process_attachment", return_value=("Summary of biochemistry chapter", True)), \
         patch("services.builder.call_llm", return_value=mock_ready_reply) as mock_llm:
        res = create_session_with_idea(
            "Help me understand cellular respiration",
            attachment=("biochemistry_handbook.pdf", b"fake long bytes"),
            db_path=test_db,
        )

        assert res["status"] == "ready"
        final_prompt = res["final_prompt"]
        assert "[PASTE FULL biochemistry_handbook.pdf HERE]" in final_prompt

        # Verify system prompt instructed placeholder
        call_system_prompt = mock_llm.call_args[0][0][0]["content"]
        assert "ATTACHED MATERIAL SUMMARY (biochemistry_handbook.pdf):" in call_system_prompt
        assert "[PASTE FULL biochemistry_handbook.pdf HERE]" in call_system_prompt


# --- API Endpoint Integration Tests ---

def test_api_create_session_with_attachment(client: FlaskClient) -> None:
    """Verify POST /api/sessions accepts multipart/form-data with file and idea."""
    mock_reply = json.dumps({
        "status": "ask",
        "type": "study",
        "question": "What level of physics is this?",
    })
    with patch("services.builder.call_llm", return_value=mock_reply):
        data = {
            "idea": "Study quantum mechanics",
            "file": (io.BytesIO(b"Wave-particle duality, Schrodinger equation, Heisenberg uncertainty."), "physics.txt"),
        }
        resp = client.post("/api/sessions", data=data, content_type="multipart/form-data")
        assert resp.status_code == 201
        res = resp.get_json()
        assert res["status"] == "ask"
        assert res["has_attachment"] is True
        assert res["attachment_filename"] == "physics.txt"
        session_id = res["session_id"]

        # Verify GET /api/sessions/<id> returns attachment fields
        get_resp = client.get(f"/api/sessions/{session_id}")
        assert get_resp.status_code == 200
        get_data = get_resp.get_json()
        assert get_data["has_attachment"] is True
        assert get_data["attachment_filename"] == "physics.txt"

        # Verify GET /api/sessions returns attachment fields in history list
        list_resp = client.get("/api/sessions")
        assert list_resp.status_code == 200
        list_data = list_resp.get_json()
        matching = [s for s in list_data if s["id"] == session_id]
        assert len(matching) == 1
        assert matching[0]["has_attachment"] is True
        assert matching[0]["attachment_filename"] == "physics.txt"


def test_api_reject_file_over_10mb(client: FlaskClient) -> None:
    """Verify that uploading a file larger than 10MB returns 400 with clear error."""
    # 10MB + 10 bytes
    large_bytes = b"x" * (10 * 1024 * 1024 + 10)
    data = {
        "idea": "Study this huge document",
        "file": (io.BytesIO(large_bytes), "huge.txt"),
    }
    resp = client.post("/api/sessions", data=data, content_type="multipart/form-data")
    assert resp.status_code == 400
    res = resp.get_json()
    assert res == {"error": "File size exceeds 10MB limit"}
