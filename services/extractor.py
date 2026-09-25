import io
from pathlib import Path
from typing import Any

from pypdf import PdfReader
from docx import Document

from services.llm import call_llm

PROMPTS_DIR = Path(__file__).resolve().parent.parent / "prompts"
MAX_EXTRACT_WORDS = 15000
SUMMARIZATION_THRESHOLD_WORDS = 3000


class ExtractionError(Exception):
    """Raised when file text extraction fails due to corruption or unsupported format."""
    pass


def count_words(text: str) -> int:
    """Count words in text using whitespace splitting."""
    return len(text.split())


def cap_words(text: str, max_words: int = MAX_EXTRACT_WORDS) -> str:
    """Cap text at roughly max_words words."""
    words = text.split()
    if len(words) > max_words:
        return " ".join(words[:max_words])
    return text


def extract_text_from_file(file_bytes: bytes, filename: str, max_words: int = MAX_EXTRACT_WORDS) -> str:
    """Extract raw text from PDF, DOCX, or TXT bytes and cap at max_words words."""
    if not file_bytes:
        raise ExtractionError("Uploaded file is empty")

    ext = Path(filename).suffix.lower()

    if ext == ".txt":
        try:
            text = file_bytes.decode("utf-8").strip()
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("latin-1", errors="replace").strip()
            except Exception as exc:
                raise ExtractionError(f"Failed to decode text file: {exc}") from exc

    elif ext == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            pages = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    pages.append(extracted)
            text = "\n".join(pages).strip()
        except Exception as exc:
            raise ExtractionError(f"Failed to extract text from PDF: {exc}") from exc

    elif ext == ".docx":
        try:
            doc = Document(io.BytesIO(file_bytes))
            parts = []
            for p in doc.paragraphs:
                if p.text and p.text.strip():
                    parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts).strip()
        except Exception as exc:
            raise ExtractionError(f"Failed to extract text from DOCX: {exc}") from exc

    else:
        raise ExtractionError(f"Unsupported file format '{ext}'. Supported formats: .pdf, .docx, .txt")

    if not text:
        raise ExtractionError("File contains no extractable text")

    return cap_words(text, max_words=max_words)


def summarize_document(text: str) -> str:
    """Summarize long text down to a few hundred words using a single LLM call."""
    template_path = PROMPTS_DIR / "summarize.txt"
    template = template_path.read_text(encoding="utf-8")
    prompt = template.replace("{text}", text)
    messages = [{"role": "user", "content": prompt}]
    return call_llm(messages, json_mode=False, max_tokens=1000)


def process_attachment(file_bytes: bytes, filename: str) -> tuple[str, bool]:
    """Extract and process attachment, summarizing if word count exceeds threshold.

    Returns tuple of (processed_text_or_summary, was_summarized).
    """
    raw_text = extract_text_from_file(file_bytes, filename)
    word_count = count_words(raw_text)

    if word_count > SUMMARIZATION_THRESHOLD_WORDS:
        summary = summarize_document(raw_text)
        return summary.strip(), True

    return raw_text, False
